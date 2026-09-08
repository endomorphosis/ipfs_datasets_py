#!/usr/bin/env python3
"""Normalize Vaquill Open US Law snapshot dumps onto NormalizedStatute.

The remaining official live scrapers (NH FortiWeb, WI docs.legis geo-block,
TN Lexis JS/anti-bot) fail from this host because we do not use residential
proxies. This ingester maps already-downloaded Vaquill v2026.08 files onto
the scraper NormalizedStatute schema for local use.

This output is a point-in-time snapshot. It is not a current-bundle, does
not authorize Hub publication, and must not be written into a live
acquisition-evidence dest.
"""

from __future__ import annotations

import argparse
import gzip
import hashlib
import io
import json
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, Iterator, List, Optional

HERE = Path(__file__).resolve()
WORKTREE = HERE.parents[3]
if str(WORKTREE) not in sys.path:
    sys.path.insert(0, str(WORKTREE))

from ipfs_datasets_py.processors.legal_scrapers.state_scrapers.base_scraper import (  # noqa: E402
    NormalizedStatute,
    StatuteMetadata,
)

SNAPSHOT_VERSION = "v2026.08"
SNAPSHOT_DATE = "2026-08-14"
SCRAPER_VERSION = "vaquill-snapshot-normalize/1.0"
SOURCE_KIND = "vaquill_open_us_law_snapshot"
LIVE_FETCH_BLOCKED_REASON = "residential_proxy_not_used"

STATE_NAMES = {
    "AR": "Arkansas",
    "GA": "Georgia",
    "MS": "Mississippi",
    "NH": "New Hampshire",
    "NY": "New York",
    "TN": "Tennessee",
    "WI": "Wisconsin",
}

CODE_NAMES = {
    "AR": "Arkansas Code Annotated",
    "GA": "Official Code of Georgia Annotated",
    "MS": "Mississippi Code of 1972",
    "NH": "New Hampshire Revised Statutes Annotated",
    "NY": "Consolidated Laws of New York",
    "TN": "Tennessee Code Annotated",
    "WI": "Wisconsin Statutes",
}

LIVE_FETCH_NOTES = {
    "AR": (
        "Arkansas official code is served through a Lexis public-access "
        "portal. This dump is a Vaquill snapshot, not a live official fetch."
    ),
    "GA": (
        "This dump was crawled from codes.findlaw.com, not the official "
        "Georgia General Assembly / Lexis public-access code. FindLaw cannot "
        "authorize a current-bundle."
    ),
    "MS": (
        "Mississippi official code is served through a Lexis public-access "
        "portal. Most dump rows have no per-section official URL."
    ),
    "NH": (
        "Official RSA HTML at gencourt.state.nh.us / gc.nh.gov is behind "
        "FortiWeb Cloud WAF and geo-restricts non-US IPs. Vaquill fetches it "
        "through a US residential/datacenter proxy. This host does not."
    ),
    "NY": (
        "Dump is generated from official Senate no-key PDFs. It still infers "
        "printed-as-current status and does not close event-conditioned "
        "residuals, so it is not a current-bundle seal."
    ),
    "TN": (
        "The official public-access portal is Lexis hottopics/tncode "
        "(JavaScript/anti-bot). Vaquill used a headless/proxy path. This host "
        "does not use residential proxies or a licensed Lexis feed."
    ),
    "WI": (
        "docs.legis.wisconsin.gov fails DNS/TCP from this host (geo-block). "
        "Vaquill routes it through a US residential proxy. This host does not."
    ),
}


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _as_str(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_breadcrumb(value: Any) -> Any:
    if isinstance(value, list):
        return value
    if isinstance(value, str):
        stripped = value.strip()
        if stripped.startswith("[") or stripped.startswith("{"):
            try:
                return json.loads(stripped)
            except json.JSONDecodeError:
                return value
    return value


def _status_fields(status: Any) -> Dict[str, Any]:
    label = _as_str(status) or ""
    lowered = label.lower()
    return {
        "repealed": lowered == "repealed",
        "reserved": lowered == "reserved",
        "in_force": lowered in {"in_force", "in force", ""},
        "label": label or None,
    }


def _normalize_record(
    *,
    state_code: str,
    rec: Dict[str, Any],
    source_archive: str,
    source_member: str,
) -> NormalizedStatute:
    code = state_code.upper()
    status = _status_fields(
        rec.get("status") or rec.get("act_status")
    )
    citation = _as_str(rec.get("citation") or rec.get("citation_short")) or ""
    act_id = (
        _as_str(rec.get("act_id") or rec.get("record_id"))
        or citation
        or f"{_as_str(rec.get('law_id')) or code}:{_as_str(rec.get('section_id') or rec.get('section_number')) or 'unknown'}"
    )
    section_title = _as_str(
        rec.get("section_title")
        or rec.get("section_name")
        or rec.get("heading")
    )
    source_url = (
        _as_str(rec.get("official_source_url"))
        or _as_str(rec.get("official_pdf_url"))
        or _as_str(rec.get("source_url"))
        or _as_str(rec.get("source_final_url"))
        or _as_str(rec.get("verification_url"))
        or ""
    )
    last_amended = rec.get("last_amended_year")
    unofficial_mirror = None
    host = ""
    if "://" in source_url:
        host = source_url.split("/", 3)[2].lower()
        if host.endswith("findlaw.com") or host.endswith("public.law"):
            unofficial_mirror = host
    return NormalizedStatute(
        state_code=code,
        state_name=STATE_NAMES[code],
        statute_id=act_id,
        code_name=CODE_NAMES[code],
        title_number=_as_str(rec.get("title_number") or rec.get("law_id")),
        title_name=_as_str(
            rec.get("title_label")
            or rec.get("title_name")
            or rec.get("law_name")
            or rec.get("source_title_name")
        ),
        chapter_number=_as_str(rec.get("chapter") or rec.get("chapter_number")),
        chapter_name=_as_str(rec.get("chapter_name")),
        section_number=_as_str(rec.get("section_number") or rec.get("section_id")),
        section_name=section_title,
        short_title=section_title,
        full_text=_as_str(rec.get("text") or rec.get("full_text")) or "",
        source_url=source_url,
        official_cite=citation or None,
        metadata=StatuteMetadata(
            last_amended=_as_str(last_amended),
            enacted_year=_as_str(rec.get("year") or rec.get("dataset_year")),
            repealed=bool(status["repealed"]),
        ),
        structured_data={
            "vaquill_act_id": act_id,
            "vaquill_snapshot": SNAPSHOT_VERSION,
            "vaquill_snapshot_date": SNAPSHOT_DATE,
            "source_kind": SOURCE_KIND,
            "source_archive": source_archive,
            "source_member": source_member,
            "authorizing_for_publication": False,
            "current_bundle": False,
            "live_official_fetch_blocked_reason": LIVE_FETCH_BLOCKED_REASON,
            "live_official_fetch_note": LIVE_FETCH_NOTES[code],
            "unofficial_mirror": unofficial_mirror,
            "status": status["label"],
            "reserved": status["reserved"],
            "breadcrumb": _parse_breadcrumb(rec.get("breadcrumb") or rec.get("context_path")),
            "display_path": rec.get("display_path"),
            "word_count": rec.get("word_count") or rec.get("clean_word_count"),
            "text_sha256": rec.get("text_sha256")
            or rec.get("clean_text_sha256")
            or rec.get("content_sha256"),
            "citation_short": rec.get("citation_short"),
            "law_type": rec.get("law_type"),
            "record_type": rec.get("record_type") or rec.get("document_type"),
        },
        scraped_at=_utc_now(),
        scraper_version=SCRAPER_VERSION,
    )


def _iter_jsonl_bytes(raw: bytes) -> Iterator[Dict[str, Any]]:
    if raw.startswith(b"\x1f\x8b"):
        raw = gzip.decompress(raw)
    for line in raw.splitlines():
        if not line.strip():
            continue
        rec = json.loads(line)
        if isinstance(rec, dict):
            yield rec


def _iter_jsonl_path(path: Path) -> Iterator[Dict[str, Any]]:
    opener = gzip.open if path.name.endswith(".gz") else path.open
    with opener(path, "rt", encoding="utf-8") as handle:  # type: ignore[arg-type]
        for line in handle:
            if not line.strip():
                continue
            rec = json.loads(line)
            if isinstance(rec, dict):
                yield rec


def _write_statutes(path: Path, statutes: Iterable[NormalizedStatute]) -> Dict[str, Any]:
    path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    empty_text = 0
    statuses: Dict[str, int] = {}
    with path.open("w", encoding="utf-8") as handle:
        for statute in statutes:
            row = statute.to_dict()
            handle.write(json.dumps(row, ensure_ascii=False, default=str) + "\n")
            count += 1
            if not str(statute.full_text or "").strip():
                empty_text += 1
            label = str((statute.structured_data or {}).get("status") or "")
            statuses[label] = statuses.get(label, 0) + 1
    return {
        "path": str(path),
        "statute_count": count,
        "empty_full_text": empty_text,
        "status_counts": statuses,
        "bytes": path.stat().st_size,
        "sha256": _sha256_file(path),
    }


def ingest_new_hampshire(source: Path, output_root: Path) -> Dict[str, Any]:
    statutes = (
        _normalize_record(
            state_code="NH",
            rec=rec,
            source_archive=source.name,
            source_member=source.name,
        )
        for rec in _iter_jsonl_path(source)
        if str(rec.get("record_type") or rec.get("document_type") or "").lower()
        in {"statute_section", "statute", ""}
    )
    stats = _write_statutes(output_root / "NH" / "statutes.jsonl", statutes)
    stats.update(
        {
            "state_code": "NH",
            "source_path": str(source),
            "source_sha256": _sha256_file(source),
        }
    )
    return stats


def ingest_wisconsin(source: Path, output_root: Path) -> Dict[str, Any]:
    with zipfile.ZipFile(source) as bundle:
        inner = bundle.read("wisconsin-law-jsonl-v2026.08.zip")
    with zipfile.ZipFile(io.BytesIO(inner)) as jsonl_zip:
        raw = jsonl_zip.read("us_wi_statutes.jsonl.gz")
    statutes = (
        _normalize_record(
            state_code="WI",
            rec=rec,
            source_archive=source.name,
            source_member="wisconsin-law-jsonl-v2026.08.zip/us_wi_statutes.jsonl.gz",
        )
        for rec in _iter_jsonl_bytes(raw)
        if str(rec.get("document_type") or "statute").lower() == "statute"
    )
    stats = _write_statutes(output_root / "WI" / "statutes.jsonl", statutes)
    stats.update(
        {
            "state_code": "WI",
            "source_path": str(source),
            "source_sha256": _sha256_file(source),
            "skipped_corpora": [
                "Wisconsin Administrative Code",
                "Wisconsin Constitution",
                "Wisconsin Court Rules",
            ],
            "skip_reason": "current-bundle schema is compiled statutes only",
        }
    )
    return stats


def _tn_workbook_rows(workbook_bytes: bytes) -> Iterator[Dict[str, Any]]:
    import openpyxl

    workbook = openpyxl.load_workbook(
        io.BytesIO(workbook_bytes),
        read_only=True,
        data_only=True,
    )
    try:
        if "Record_Index" not in workbook.sheetnames or "Law_Text" not in workbook.sheetnames:
            raise RuntimeError("Tennessee workbook missing Record_Index/Law_Text")
        index_rows: Dict[str, Dict[str, Any]] = {}
        index_sheet = workbook["Record_Index"]
        index_iter = index_sheet.iter_rows(values_only=True)
        index_header = [str(col or "").strip() for col in next(index_iter)]
        for raw in index_iter:
            row = {
                index_header[i]: raw[i] if i < len(raw) else None
                for i in range(len(index_header))
            }
            if str(row.get("Corpus") or "").strip() != "Statute":
                continue
            record_id = _as_str(row.get("Record ID"))
            if not record_id:
                continue
            index_rows[record_id] = row

        chunks: Dict[str, List[tuple[int, str]]] = {}
        text_sheet = workbook["Law_Text"]
        text_iter = text_sheet.iter_rows(values_only=True)
        text_header = [str(col or "").strip() for col in next(text_iter)]
        for raw in text_iter:
            row = {
                text_header[i]: raw[i] if i < len(raw) else None
                for i in range(len(text_header))
            }
            if str(row.get("Corpus") or "").strip() != "Statute":
                continue
            record_id = _as_str(row.get("Record ID"))
            if not record_id:
                continue
            part = int(row.get("Chunk Part") or 1)
            text = _as_str(row.get("Text")) or ""
            chunks.setdefault(record_id, []).append((part, text))
    finally:
        workbook.close()

    missing_index = sorted(set(chunks) - set(index_rows))
    if missing_index:
        raise RuntimeError(
            "Tennessee Law_Text records missing Record_Index rows: "
            + ", ".join(missing_index[:8])
        )
    for record_id, meta in index_rows.items():
        ordered = sorted(chunks.get(record_id, []), key=lambda item: item[0])
        yield {
            "act_id": record_id,
            "citation": meta.get("Citation"),
            "title_number": meta.get("Title Number"),
            "title_name": meta.get("Title Name"),
            "chapter": meta.get("Chapter"),
            "section_number": meta.get("Section Number"),
            "section_title": meta.get("Section Title"),
            "display_path": meta.get("Display Path"),
            "act_status": meta.get("Status"),
            "text": "".join(text for _, text in ordered),
            "word_count": meta.get("Clean Word Count"),
            "source_url": meta.get("Verification URL"),
            "official_source_url": meta.get("Verification URL"),
            "text_sha256": meta.get("Clean Text SHA-256"),
            "document_type": "statute",
            "state": "TN",
        }


def ingest_tennessee(source: Path, output_root: Path) -> Dict[str, Any]:
    statutes: List[NormalizedStatute] = []
    members: List[str] = []
    with zipfile.ZipFile(source) as bundle:
        names = [
            name
            for name in bundle.namelist()
            if name.startswith("tennessee_code_titles_") and name.endswith(".xlsx")
        ]
        if not names:
            raise RuntimeError("Tennessee zip has no title workbooks")
        for name in sorted(names):
            members.append(name)
            for rec in _tn_workbook_rows(bundle.read(name)):
                statutes.append(
                    _normalize_record(
                        state_code="TN",
                        rec=rec,
                        source_archive=source.name,
                        source_member=name,
                    )
                )
    stats = _write_statutes(output_root / "TN" / "statutes.jsonl", statutes)
    stats.update(
        {
            "state_code": "TN",
            "source_path": str(source),
            "source_sha256": _sha256_file(source),
            "source_members": members,
            "skipped_corpora": [
                "Tennessee Constitution",
                "Tennessee court rules",
                "Tennessee agency guidance",
            ],
            "skip_reason": "current-bundle schema is compiled statutes only",
            "per_section_official_url": False,
            "source_url_note": (
                "Vaquill Excel Verification URL is the Lexis portal homepage, "
                "not a per-section locator."
            ),
        }
    )
    return stats


def _is_compiled_statute_row(rec: Dict[str, Any]) -> bool:
    kind = str(
        rec.get("record_type") or rec.get("document_type") or ""
    ).strip().lower()
    return kind in {"statute", "statute_section", "section"}


def ingest_arkansas(source: Path, output_root: Path) -> Dict[str, Any]:
    statutes = (
        _normalize_record(
            state_code="AR",
            rec=rec,
            source_archive=source.name,
            source_member=source.name,
        )
        for rec in _iter_jsonl_path(source)
        if _is_compiled_statute_row(rec)
        and str(rec.get("corpus") or rec.get("normalized_corpus") or "")
        .casefold()
        in {"arkansas code", "statute", ""}
    )
    stats = _write_statutes(output_root / "AR" / "statutes.jsonl", statutes)
    stats.update(
        {
            "state_code": "AR",
            "source_path": str(source),
            "source_sha256": _sha256_file(source),
            "skipped_corpora": ["Arkansas Constitution"],
            "skip_reason": "current-bundle schema is compiled statutes only",
        }
    )
    return stats


def ingest_georgia(source: Path, output_root: Path) -> Dict[str, Any]:
    statutes = (
        _normalize_record(
            state_code="GA",
            rec=rec,
            source_archive=source.name,
            source_member=source.name,
        )
        for rec in _iter_jsonl_path(source)
        if _is_compiled_statute_row(rec)
    )
    stats = _write_statutes(output_root / "GA" / "statutes.jsonl", statutes)
    stats.update(
        {
            "state_code": "GA",
            "source_path": str(source),
            "source_sha256": _sha256_file(source),
            "skipped_corpora": ["Georgia Constitution"],
            "skip_reason": "current-bundle schema is compiled statutes only",
            "unofficial_mirror": "codes.findlaw.com",
            "source_url_note": (
                "FindLaw cannot authorize a current-bundle. Rows are retained "
                "as a non-authorizing snapshot only."
            ),
        }
    )
    return stats


def ingest_mississippi(source: Path, output_root: Path) -> Dict[str, Any]:
    statutes = (
        _normalize_record(
            state_code="MS",
            rec=rec,
            source_archive=source.name,
            source_member=source.name,
        )
        for rec in _iter_jsonl_path(source)
        if _is_compiled_statute_row(rec)
        and str(rec.get("corpus_name") or "").casefold()
        in {"mississippi code of 1972", ""}
    )
    stats = _write_statutes(output_root / "MS" / "statutes.jsonl", statutes)
    stats.update(
        {
            "state_code": "MS",
            "source_path": str(source),
            "source_sha256": _sha256_file(source),
            "skipped_corpora": [
                "Mississippi Constitution",
                "Mississippi Court Rules",
            ],
            "skip_reason": "current-bundle schema is compiled statutes only",
        }
    )
    return stats


def ingest_new_york(source: Path, output_root: Path) -> Dict[str, Any]:
    statutes = (
        _normalize_record(
            state_code="NY",
            rec=rec,
            source_archive=source.name,
            source_member=source.name,
        )
        for rec in _iter_jsonl_path(source)
        if str(rec.get("record_type") or "").strip().lower() == "section"
        and str(rec.get("law_type") or "").strip().upper() == "CONSOLIDATED"
    )
    stats = _write_statutes(output_root / "NY" / "statutes.jsonl", statutes)
    stats.update(
        {
            "state_code": "NY",
            "source_path": str(source),
            "source_sha256": _sha256_file(source),
            "skipped_corpora": [
                "front_matter",
                "legislative rules",
                "court acts",
                "unconsolidated laws",
                "constitution",
            ],
            "skip_reason": "current-bundle schema is consolidated statutes only",
            "source_url_note": (
                "Official Senate PDF snapshot without event-proof closure; "
                "not a current-bundle seal."
            ),
        }
    )
    return stats


def _write_readme(output_root: Path, manifest: Dict[str, Any]) -> None:
    lines = [
        "# Vaquill snapshot normalized to NormalizedStatute",
        "",
        f"Snapshot: {SNAPSHOT_VERSION} ({SNAPSHOT_DATE})",
        "Schema: `ipfs_datasets_py.processors.legal_scrapers.state_scrapers.base_scraper.NormalizedStatute`",
        "",
        "This is a point-in-time research dump. It is **not current law**, is",
        "**not a current-bundle**, and is **not authorizing for Hub publication**.",
        "",
        "## Why this path exists",
        "",
        "Official live scrapers for the remaining states fail from this host",
        "because we do not use residential proxies (or a licensed Lexis feed):",
        "",
        "- NH: FortiWeb geo-WAF on gencourt/gc.nh.gov",
        "- WI: docs.legis.wisconsin.gov DNS/TCP geo-block",
        "- TN/AR/MS: Lexis public-access portals",
        "- GA dump is FindLaw, not official",
        "- NY dump is official Senate PDFs without event-proof closure",
        "",
        "Vaquill's remaining-state fetch stack is US proxy + optional headless.",
        "These files are already-scraped snapshots remapped onto the schema",
        "our adapters emit.",
        "",
        "## Outputs",
        "",
    ]
    for state in manifest.get("states", []):
        lines.append(
            f"- `{state['state_code']}/statutes.jsonl`: {state['statute_count']} statutes"
        )
    lines.extend(
        [
            "",
            "This tree must not be seeded as a current-bundle acquisition-evidence dest.",
            "",
        ]
    )
    (output_root / "README.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input-root",
        type=Path,
        default=Path("/tmp/tmp_laws"),
        help="Directory containing the downloaded Vaquill archives",
    )
    parser.add_argument(
        "--output-root",
        type=Path,
        default=Path.home()
        / ".ipfs_datasets"
        / "state_laws"
        / "vaquill-snapshot-normalized-v2026.08.31",
        help="Non-authorizing NormalizedStatute JSONL output root",
    )
    args = parser.parse_args()

    output_root = args.output_root.resolve()
    forbidden = (
        "acquisition-evidence",
        "live-v",
        "current-live",
        "permanently-nonauthorizing",
    )
    if any(token in str(output_root) for token in forbidden):
        raise SystemExit(f"REFUSING output root looks like a live dest: {output_root}")
    if output_root.exists() and any(output_root.iterdir()):
        raise SystemExit(f"REFUSING output root already exists: {output_root}")

    input_root = args.input_root.resolve()
    sources = {
        "AR": input_root / "Arkansas_Law_v2026-08.cleaned.jsonl.gz",
        "GA": input_root / "Georgia_Laws_FRESH_2026-08-31.jsonl.gz",
        "MS": input_root / "Mississippi_Laws_2026-08-31.jsonl.gz",
        "NH": input_root / "New_Hampshire_Revised_Statutes_v2026-08.jsonl.gz",
        "NY": input_root / "New_York_Laws_2026-08-31.jsonl.gz",
        "TN": input_root / "Tennessee_Laws_Complete_Collection_REISSUED_2026-08.zip",
        "WI": input_root / "Wisconsin_Legal_Corpus_2026-08-31_bundle.zip",
    }
    missing = [code for code, path in sources.items() if not path.is_file()]
    if missing:
        raise SystemExit(f"missing source archives for {missing} under {input_root}")

    output_root.mkdir(parents=True, exist_ok=True)
    states = [
        ingest_arkansas(sources["AR"], output_root),
        ingest_georgia(sources["GA"], output_root),
        ingest_mississippi(sources["MS"], output_root),
        ingest_new_hampshire(sources["NH"], output_root),
        ingest_new_york(sources["NY"], output_root),
        ingest_tennessee(sources["TN"], output_root),
        ingest_wisconsin(sources["WI"], output_root),
    ]
    manifest = {
        "schema": "ipfs_datasets_py.state_laws.vaquill_snapshot_normalized.v1",
        "authorizing_for_publication": False,
        "current_bundle": False,
        "live_official_fetch_blocked_reason": LIVE_FETCH_BLOCKED_REASON,
        "vaquill_snapshot": SNAPSHOT_VERSION,
        "vaquill_snapshot_date": SNAPSHOT_DATE,
        "normalized_at": _utc_now(),
        "scraper_version": SCRAPER_VERSION,
        "normalized_schema": "NormalizedStatute",
        "input_root": str(input_root),
        "output_root": str(output_root),
        "states": states,
        "not_ingested": [],
        "notes": [
            "Point-in-time snapshot dump, not current law.",
            "Official live acquisition still fails here without residential proxies.",
            "Georgia rows are FindLaw, not official.",
            "Do not seed this tree as a current-bundle acquisition-evidence dest.",
            "This does not complete LCR exact-51 current-bundle seals.",
        ],
    }
    (output_root / "ingest_manifest.json").write_text(
        json.dumps(manifest, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    _write_readme(output_root, manifest)
    print(json.dumps(manifest, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
